"""Lectura y escritura de archivos Excel (.xlsx) para importar/exportar
clientes, productos, facturas y cotizaciones.

Formato de trabajo: la primera fila contiene los encabezados (nombres simples,
sin acentos ni espacios) y las filas siguientes los datos. Es el mismo formato
que se genera al exportar, de modo que un archivo exportado puede volver a
importarse.
"""
import os
import unicodedata
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="0F172A")
ALT_FILL = PatternFill("solid", fgColor="F1F5F9")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INSTRUCTIONS_SHEET = "Instrucciones"


def normalize_header(value) -> str:
    """Convierte un encabezado en una clave simple: minúsculas, sin acentos."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.strip().lower().replace(" ", "_").replace("-", "_")
    return text


def _auto_width(worksheet, headers, rows, max_width=44):
    for col_index, header in enumerate(headers, start=1):
        width = len(str(header)) + 2
        for row in rows:
            value = row.get(header, "")
            if value is None:
                continue
            width = max(width, min(len(str(value)) + 2, max_width))
        worksheet.column_dimensions[get_column_letter(col_index)].width = width


def _style_header(worksheet, headers):
    for col_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    worksheet.row_dimensions[1].height = 22
    worksheet.freeze_panes = "A2"


def write_table(path_or_buffer, sheet_name: str, headers, rows,
                instructions=None, autofilter: bool = True):
    """Escribe una tabla con encabezados y filas (lista de dicts).

    `instructions` es una lista de tuplas (campo, descripción, ejemplo) que se
    agrega en una segunda hoja para guiar al usuario.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    _style_header(sheet, headers)
    for row_index, row in enumerate(rows, start=2):
        for col_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=row.get(header))
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index % 2 == 0:
                cell.fill = ALT_FILL

    if autofilter and rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    _auto_width(sheet, headers, rows)

    if instructions:
        info = workbook.create_sheet(INSTRUCTIONS_SHEET)
        info_headers = ["campo", "descripcion", "ejemplo"]
        _style_header(info, info_headers)
        for row_index, (field_name, description, example) in enumerate(instructions, start=2):
            for col_index, value in enumerate((field_name, description, example), start=1):
                cell = info.cell(row=row_index, column=col_index, value=value)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        info.column_dimensions["A"].width = 26
        info.column_dimensions["B"].width = 70
        info.column_dimensions["C"].width = 26

    if isinstance(path_or_buffer, (str, os.PathLike)):
        os.makedirs(os.path.dirname(os.path.abspath(path_or_buffer)), exist_ok=True)
        workbook.save(path_or_buffer)
        return str(path_or_buffer)

    workbook.save(path_or_buffer)
    return path_or_buffer


def table_bytes(sheet_name: str, headers, rows, instructions=None) -> bytes:
    buffer = BytesIO()
    write_table(buffer, sheet_name, headers, rows, instructions=instructions)
    return buffer.getvalue()


def read_table(path_or_file, sheet_name: str = None):
    """Lee un .xlsx y devuelve (encabezados, filas_como_dicts).

    La primera fila con contenido se toma como encabezado. Se ignoran las filas
    totalmente vacías.
    """
    workbook = load_workbook(path_or_file, read_only=True, data_only=True)

    if sheet_name and sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
    else:
        # Prefiere una hoja de datos (no la de instrucciones).
        candidates = [s for s in workbook.sheetnames if s != INSTRUCTIONS_SHEET]
        sheet = workbook[candidates[0]] if candidates else workbook.active

    raw_rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    headers = None
    data = []
    for raw in raw_rows:
        if raw is None:
            continue
        values = list(raw)
        if headers is None:
            if not any(v is not None and str(v).strip() != "" for v in values):
                continue
            headers = [normalize_header(v) for v in values]
            continue
        if not any(v is not None and str(v).strip() != "" for v in values):
            continue
        row = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            row[header] = values[index] if index < len(values) else None
        data.append(row)

    return headers or [], data
