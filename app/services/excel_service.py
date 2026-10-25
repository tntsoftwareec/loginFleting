"""
Servicio de exportación a Excel para listados del sistema.
"""
import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.paths import GENERATED_DIR
from app.services.sri_catalogs import etiqueta_forma_pago

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
CELL_ALIGN = Alignment(horizontal="left", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


class ExcelService:
    """Servicio para exportar listados a archivos Excel."""

    @classmethod
    def _create_workbook(cls, headers, rows, sheet_name="Datos"):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Escribir encabezados
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = THIN_BORDER

        # Escribir datos
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = CELL_ALIGN
                cell.border = THIN_BORDER

        # Auto-ajustar ancho de columnas
        for col_idx, header in enumerate(headers, 1):
            max_length = len(str(header))
            for row_data in rows:
                cell_value = str(row_data[col_idx - 1]) if col_idx - 1 < len(row_data) else ""
                max_length = max(max_length, len(cell_value))
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

        return wb

    @classmethod
    def _save_workbook(cls, wb, filepath):
        wb.save(filepath)
        return str(filepath)

    @classmethod
    def export_clientes(cls, clientes):
        """Exporta listado de clientes a Excel."""
        headers = ["ID", "Identificación", "Razón Social", "Tipo Documento", "Email", "Teléfono", "Dirección", "Estado"]
        rows = []
        for c in clientes:
            rows.append([
                c.get("id", ""),
                c.get("identificacion", ""),
                c.get("razon_social", ""),
                c.get("tipo_documento", ""),
                c.get("email", ""),
                c.get("telefono", ""),
                c.get("direccion", ""),
                c.get("status", ""),
            ])
        wb = cls._create_workbook(headers, rows, "Clientes")
        filepath = GENERATED_DIR / "exports" / f"clientes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath.parent.mkdir(parents=True, exist_ok=True)
        return cls._save_workbook(wb, filepath)

    @classmethod
    def export_facturas(cls, facturas):
        """Exporta listado de facturas a Excel."""
        headers = [
            "N° Factura", "Fecha Emisión", "Cliente", "Identificación",
            "Forma Pago", "Subtotal", "IVA", "ICE", "Propina", "Total", "Estado",
        ]
        rows = []
        for f in facturas:
            rows.append([
                f.get("secuencial", ""),
                f.get("fecha_emision", ""),
                f.get("client_nombre", ""),
                f.get("client_identificacion", ""),
                etiqueta_forma_pago(f.get("forma_pago", "")),
                f.get("subtotal_sin_impuestos", 0),
                f.get("iva_total", 0),
                f.get("ice_total", 0),
                f.get("propina", 0),
                f.get("importe_total", 0),
                f.get("estado", ""),
            ])
        wb = cls._create_workbook(headers, rows, "Facturas")
        filepath = GENERATED_DIR / "exports" / f"facturas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath.parent.mkdir(parents=True, exist_ok=True)
        return cls._save_workbook(wb, filepath)

    @classmethod
    def export_cotizaciones(cls, cotizaciones):
        """Exporta listado de cotizaciones a Excel."""
        headers = ["N° Cotización", "Fecha Emisión", "Válido Hasta", "Cliente", "Identificación", "Subtotal", "IVA", "Total", "Estado"]
        rows = []
        for c in cotizaciones:
            rows.append([
                c.get("secuencial", ""),
                c.get("fecha_emision", ""),
                c.get("fecha_validez", ""),
                c.get("client_nombre", ""),
                c.get("client_identificacion", ""),
                c.get("subtotal_sin_impuestos", 0),
                c.get("iva_total", 0),
                c.get("importe_total", 0),
                c.get("estado", ""),
            ])
        wb = cls._create_workbook(headers, rows, "Cotizaciones")
        filepath = GENERATED_DIR / "exports" / f"cotizaciones_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath.parent.mkdir(parents=True, exist_ok=True)
        return cls._save_workbook(wb, filepath)

    @classmethod
    def export_productos(cls, productos):
        """Exporta listado de productos a Excel."""
        headers = ["Código", "Nombre", "Descripción", "Precio", "Costo", "Stock", "IVA %", "Categoría", "Estado"]
        rows = []
        for p in productos:
            rows.append([
                p.get("codigo", ""),
                p.get("nombre", ""),
                p.get("descripcion", ""),
                p.get("precio", 0),
                p.get("costo", 0),
                p.get("stock", 0),
                p.get("iva_porcentaje", 15),
                p.get("categoria", ""),
                p.get("status", ""),
            ])
        wb = cls._create_workbook(headers, rows, "Productos")
        filepath = GENERATED_DIR / "exports" / f"productos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath.parent.mkdir(parents=True, exist_ok=True)
        return cls._save_workbook(wb, filepath)

    @classmethod
    def export_proveedores(cls, proveedores):
        """Exporta listado de proveedores a Excel."""
        headers = ["ID", "Identificación", "Razón Social", "Email", "Teléfono", "Contacto", "Dirección", "Estado"]
        rows = []
        for p in proveedores:
            rows.append([
                p.get("id", ""),
                p.get("identificacion", ""),
                p.get("razon_social", ""),
                p.get("email", ""),
                p.get("telefono", ""),
                p.get("contacto", ""),
                p.get("direccion", ""),
                p.get("status", ""),
            ])
        wb = cls._create_workbook(headers, rows, "Proveedores")
        filepath = GENERATED_DIR / "exports" / f"proveedores_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath.parent.mkdir(parents=True, exist_ok=True)
        return cls._save_workbook(wb, filepath)
